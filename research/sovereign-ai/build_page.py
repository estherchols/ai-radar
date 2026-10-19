"""Render research/sovereign-ai/index.html from the recalculated model."""
import json, re
from openpyxl import load_workbook

HERE = "/home/claude/ai-radar/research/sovereign-ai/"
wb = load_workbook(HERE + "sovereign-ai-model.xlsx", data_only=True)
REF = json.load(open(HERE + "model_refs.json"))

def by_code(sheet, code, col=3):
    for r in wb[sheet].iter_rows(min_row=5):
        if r[0].value == code:
            return r[col - 1].value
    raise KeyError(f"{sheet}:{code}")
def cell(ref):
    sheet, addr = ref.split("!")
    return wb[sheet][addr.replace("$", "")].value

V = {}
for k in ["A3", "A5", "A7", "A8", "B1", "B4", "B5", "B7", "C1", "C2", "C3", "C4", "C5", "C7", "R1", "R2", "R3", "R4", "R5", "R6", "R7", "R8", "R9", "R10", "R11"]:
    V["td_" + k] = by_code("TopDown", k)
bu = wb["BottomUp"]
progs = []
for r in bu.iter_rows(min_row=REF["first"], max_row=REF["last"]):
    progs.append(dict(name=r[0].value, quad=r[1].value, ctry=r[2].value, mw=r[3].value, yr=r[4].value, mw26=r[5].value, gpus=r[6].value, gpus26=r[7].value, capex26=r[8].value, gpusys=r[9].value, pipe=r[10].value))
tot = [c.value for c in bu[REF["tot"]]]
quads = []
for r in bu.iter_rows(min_row=REF["qstart"], max_row=REF["qstart"] + 3):
    quads.append(dict(name=r[0].value, mw26=r[1].value, gpus26=r[2].value, capex26=r[3].value, gpusys=r[4].value, mw=r[5].value))
V["gulf_share"] = bu.cell(row=REF["qstart"] + 4, column=4).value
V["eu_ann_share"] = bu.cell(row=REF["qstart"] + 5, column=6).value
V["eu_en_share"] = bu.cell(row=REF["qstart"] + 6, column=2).value
for k, ref in REF["BU"].items():
    V["bu_" + k] = cell(ref)
for k, ref in REF["GT"].items():
    V["gt_" + k] = cell(ref)
gt = wb["GigawattTCO"]
def grow_vals(code):
    for r in gt.iter_rows(min_row=5):
        if r[0].value == code:
            return [r[c].value for c in range(2, 7)]
G = {c: grow_vals(c) for c in ["G7", "G11", "G13", "G14", "G17", "G18", "G18c", "G21", "G15"]}
kc = wb["Configs"]
def krow_vals(r_):
    return [kc.cell(row=r_, column=c).value for c in (3, 4, 5, 6)]
KV = {code: krow_vals(r_) for code, r_ in REF["KC"].items()}
for k, ref in REF["PR"].items():
    V["pr_" + k] = cell("Pricing!" + ref) if "!" not in ref else cell(ref)
for k in ["P1", "P2", "P3", "P4", "P5", "P5b", "P5c", "P5d", "P6", "P7", "P8", "P9", "P10", "P11", "P12", "P13", "P14", "P15", "P16", "P17", "P18", "P19", "P20", "P21", "P22", "P23", "P24", "P25", "P25b", "P25c", "P26", "P27", "P28", "P29"]:
    V["p_" + k] = by_code("Pricing", k)
for k, ref in REF["VV"].items():
    V["vv_" + k] = cell("Value!" + ref) if "!" not in ref else cell(ref)
for k in ["VA3", "VA4", "VA5", "VA6", "VA7", "VA8", "VA9", "VA10", "VA11", "VA12", "VA13", "VA14", "VB1", "VB2", "VB3", "VC1", "VC2", "VC3", "VC3b", "VC4", "VC5", "VC6", "VC6b", "VC7"]:
    V["v_" + k] = by_code("Value", k)

def bn(x, d=1): return f"${x:,.{d}f}B"
def usd(x, d=2): return f"${x:,.{d}f}"
def pct(x, d=0): return f"{x*100:,.{d}f}%"
def num(x, d=0): return f"{x:,.{d}f}"
def gw(x, d=1): return f"{x:,.{d}f} GW"

INK = "var(--ink)"; DIM = "var(--dim)"; LINE = "var(--line)"; ACC = "var(--electric)"; ACC2 = "#e3a248"; ACC3 = "#c4561c"; GREY = "var(--dim)"
SANS = "Schibsted Grotesk,sans-serif"; MONO = "IBM Plex Mono,monospace"

def chart_sizing():
    bars = [("Path A\nNVIDIA revenue", V["td_R1"]), ("Path B\nanalysts", V["td_R2"]), ("Path C\nbudgets x leverage", V["td_R3"]), ("Top-down\naverage", V["td_R4"]), ("Bottom-up\nenergised 2026", V["td_R6"])]
    W, H = 720, 300; L, R, T, B = 50, 20, 20, 60; mx = 90
    pw = W - L - R; ph = H - T - B; bw = pw / len(bars) * 0.6
    out = [f'<svg viewBox="0 0 {W} {H}" xmlns="http://www.w3.org/2000/svg" role="img" aria-label="Top-down versus bottom-up sizing of sovereign AI compute spend 2026">']
    for g in range(0, 91, 30):
        y = T + ph - ph * g / mx
        out.append(f'<line x1="{L}" y1="{y:.1f}" x2="{W-R}" y2="{y:.1f}" stroke="{LINE}"/>')
        out.append(f'<text x="{L-8}" y="{y+4:.1f}" font-size="11" text-anchor="end" fill="{DIM}" font-family="{MONO}">${g}B</text>')
    for i, (lab, val) in enumerate(bars):
        x = L + pw / len(bars) * (i + 0.2); h = ph * val / mx; y = T + ph - h
        fill = ACC if i < 3 else (ACC2 if i == 3 else ACC3)
        out.append(f'<rect x="{x:.1f}" y="{y:.1f}" width="{bw:.1f}" height="{h:.1f}" rx="3" fill="{fill}"/>')
        out.append(f'<text x="{x+bw/2:.1f}" y="{y-6:.1f}" font-size="12" text-anchor="middle" fill="{INK}" font-weight="600" font-family="{SANS}">${val:,.0f}B</text>')
        for j, line in enumerate(lab.split("\n")):
            out.append(f'<text x="{x+bw/2:.1f}" y="{T+ph+16+j*13}" font-size="11" text-anchor="middle" fill="{DIM}" font-family="{SANS}">{line}</text>')
    out.append("</svg>"); return "\n".join(out)

def chart_pipeline():
    rows = sorted(progs, key=lambda p: -p["mw"])[:12]
    W = 720; rowh = 30; L = 300; R = 80; T = 24
    H = T + rowh * len(rows) + 10; mx = max(p["mw"] for p in rows)
    out = [f'<svg viewBox="0 0 {W} {H}" xmlns="http://www.w3.org/2000/svg" role="img" aria-label="Announced versus energised megawatts by sovereign AI programme">']
    out.append(f'<text x="{L}" y="12" font-size="11" fill="{DIM}" font-family="{MONO}">announced MW (grey) vs energised by end-2026 (teal)</text>')
    for i, p in enumerate(rows):
        y = T + i * rowh
        label = p["name"].split(":")[0].split(" (")[0][:34]
        out.append(f'<text x="{L-10}" y="{y+19}" font-size="12" text-anchor="end" fill="{INK}" font-family="{SANS}">{label}</text>')
        w1 = (W - L - R) * p["mw"] / mx; w2 = (W - L - R) * p["mw26"] / mx
        out.append(f'<rect x="{L}" y="{y+6}" width="{w1:.1f}" height="18" rx="3" fill="{LINE}"/>')
        out.append(f'<rect x="{L}" y="{y+6}" width="{max(w2,1):.1f}" height="18" rx="3" fill="{ACC}"/>')
        out.append(f'<text x="{L+w1+8:.1f}" y="{y+19}" font-size="11" fill="{DIM}" font-family="{MONO}">{p["mw"]:,.0f} / {p["mw26"]:,.0f}</text>')
    out.append("</svg>"); return "\n".join(out)

def chart_tco():
    names = ["United States", "Saudi Arabia", "South Korea", "Germany", "India"]
    W, H = 720, 320; L, R, T, B = 55, 20, 20, 50
    pw = W - L - R; ph = H - T - B; ymax = 12
    def Y(v): return T + ph - ph * v / ymax
    out = [f'<svg viewBox="0 0 {W} {H}" xmlns="http://www.w3.org/2000/svg" role="img" aria-label="Owner cost per GPU hour by country against rental prices">']
    for g in range(0, 13, 2):
        out.append(f'<line x1="{L}" y1="{Y(g):.1f}" x2="{W-R}" y2="{Y(g):.1f}" stroke="{LINE}"/>')
        out.append(f'<text x="{L-8}" y="{Y(g)+4:.1f}" font-size="11" text-anchor="end" fill="{DIM}" font-family="{MONO}">${g}</text>')
    gw_ = pw / len(names); bw = gw_ * 0.28
    for i, n in enumerate(names):
        x0 = L + gw_ * i + gw_ * 0.15
        for j, (vals, col, lab) in enumerate([(G["G18"], ACC, "cash cost incl. delay"), (G["G18c"], ACC2, "incl. 12% cost of capital")]):
            v = vals[i]; x = x0 + j * (bw + 4)
            out.append(f'<rect x="{x:.1f}" y="{Y(v):.1f}" width="{bw:.1f}" height="{ph*v/ymax:.1f}" rx="3" fill="{col}"/>')
            out.append(f'<text x="{x+bw/2:.1f}" y="{Y(v)-5:.1f}" font-size="10.5" text-anchor="middle" fill="{INK}" font-family="{MONO}">${v:.2f}</text>')
        out.append(f'<text x="{x0+bw+2:.1f}" y="{T+ph+18}" font-size="11" text-anchor="middle" fill="{DIM}" font-family="{SANS}">{n}</text>')
    for price, col, lab in [(V["p_P6"], ACC3, "CoreWeave GB200 on-demand $10.50"), (V["p_P7"], ACC3, "Nebius B200 $8.50"), (V["p_P15"], GREY, "Scaleway H100 (Hopper) $3.10")]:
        out.append(f'<line x1="{L}" y1="{Y(price):.1f}" x2="{W-R}" y2="{Y(price):.1f}" stroke="{col}" stroke-width="1.5" stroke-dasharray="6 4"/>')
        out.append(f'<text x="{W-R-4}" y="{Y(price)-5:.1f}" font-size="11" text-anchor="end" fill="{col}" font-family="{SANS}">{lab}</text>')
    out.append(f'<rect x="{L}" y="{H-14}" width="10" height="10" fill="{ACC}"/><text x="{L+14}" y="{H-5}" font-size="10.5" fill="{DIM}" font-family="{SANS}">owner cash cost incl. delay</text>')
    out.append(f'<rect x="{L+200}" y="{H-14}" width="10" height="10" fill="{ACC2}"/><text x="{L+214}" y="{H-5}" font-size="10.5" fill="{DIM}" font-family="{SANS}">incl. 12% cost of capital</text>')
    out.append("</svg>"); return "\n".join(out)

def chart_matrix():
    W, H = 720, 520; L, T = 70, 30; pw, ph = 600, 420
    out = [f'<svg viewBox="0 0 {W} {H}" xmlns="http://www.w3.org/2000/svg" role="img" aria-label="Two by two map of sovereign AI buyers">']
    out.append(f'<rect x="{L}" y="{T}" width="{pw}" height="{ph}" fill="none" stroke="{LINE}"/>')
    out.append(f'<line x1="{L+pw/2}" y1="{T}" x2="{L+pw/2}" y2="{T+ph}" stroke="{LINE}"/><line x1="{L}" y1="{T+ph/2}" x2="{L+pw}" y2="{T+ph/2}" stroke="{LINE}"/>')
    out.append(f'<text x="{L+pw/2}" y="{T+ph+28}" font-size="12" text-anchor="middle" fill="{DIM}" font-family="{SANS}">Power: time and price to energise a gigawatt  →  (fast and cheap on the right)</text>')
    out.append(f'<text transform="translate(22,{T+ph/2}) rotate(-90)" font-size="12" text-anchor="middle" fill="{DIM}" font-family="{SANS}">Capital: depth of state and sovereign-wealth balance sheet  →</text>')
    q = {x["name"]: x for x in quads}
    qd = [
        (L + pw / 2, T, "Compute exporters", "Saudi Arabia · UAE · Qatar", "Sovereign wealth plus cheap gas and solar. Build beyond domestic need and sell capacity to foreign labs. Chips negotiated at head-of-state level.", q["Compute exporters"], ACC3),
        (L, T, "Industrial sovereigns", "Japan · South Korea · Taiwan", "Deep corporate balance sheets, constrained grids and land. Chaebol and telco AI factories for the group's own cars, chips and robots. First in line for Rubin.", q["Industrial sovereigns"], ACC2),
        (L, T + ph / 2, "Regulated builders", "EU · France · Germany · UK", "Budgets and development banks, slow permits and the dearest power in the set. Announce the most megawatts, energise the fewest. Research tools, not clouds.", q["Regulated builders"], GREY),
        (L + pw / 2, T + ph / 2, "Subsidised tenants", "India · Malaysia · Thailand · LatAm", "Thin budgets, decent power. The state subsidises the GPU-hour and lets hyperscalers and neoclouds build. Cheapest way to put compute in startups' hands.", q["Subsidised tenants"], ACC),
    ]
    def wrap(text, n):
        words = text.split(); lines = []; cur = ""
        for w in words:
            if len(cur) + len(w) + 1 > n and cur: lines.append(cur); cur = w
            else: cur = (cur + " " + w).strip()
        lines.append(cur); return lines
    for x, y, t, who, body, qq, col in qd:
        out.append(f'<rect x="{x+10}" y="{y+10}" width="{pw/2-20}" height="{ph/2-20}" rx="10" fill="var(--bg)" stroke="{col}" stroke-width="1.5"/>')
        out.append(f'<text x="{x+24}" y="{y+40}" font-size="17" font-weight="700" fill="{INK}" font-family="{SANS}">{t}</text>')
        out.append(f'<text x="{x+pw/2-24}" y="{y+60}" font-size="11" text-anchor="end" fill="{col}" font-family="{MONO}">{qq["mw26"]:,.0f} MW live 2026</text>')
        yy = y + 60
        items = [i.strip() for i in who.split("·")]; wl=[]; cur=""
        for it in items:
            cand = (cur + " · " + it) if cur else it
            if len(cand) > 20 and cur: wl.append(cur); cur = it
            else: cur = cand
        wl.append(cur)
        for ln in wl:
            out.append(f'<text x="{x+24}" y="{yy}" font-size="11" fill="{col}" font-family="{MONO}">{ln}</text>'); yy += 15
        yy += 6
        for ln in wrap(body, 38)[:8]:
            out.append(f'<text x="{x+24}" y="{yy}" font-size="12.5" fill="{INK}" font-family="{SANS}">{ln}</text>'); yy += 17
        out.append(f'<text x="{x+24}" y="{y+ph/2-22}" font-size="11" fill="{DIM}" font-family="{MONO}">{qq["mw"]:,.0f} MW announced · ${qq["capex26"]/1000:,.1f}B capex 2026</text>')
    out.append("</svg>"); return "\n".join(out)

def chart_margins():
    rows = [("Saudi exporter campus, after 12% capital (VC3b)", V["v_VC3b"]), ("Saudi exporter campus, cash margin (VC3)", V["v_VC3"]), ("EU gigafactory at European cloud prices, cash (VC6)", V["v_VC6"]), ("EU gigafactory at European cloud prices, after 12% capital (VC6b)", V["v_VC6b"])]
    W = 720; rowh = 34; L = 360; T = 10; H = T + rowh * len(rows) + 10
    mid = L + (W - L - 20) / 2; span = (W - L - 20) / 2
    out = [f'<svg viewBox="0 0 {W} {H}" xmlns="http://www.w3.org/2000/svg" role="img" aria-label="Operating margin of a sovereign gigawatt by quadrant">']
    out.append(f'<line x1="{mid}" y1="{T}" x2="{mid}" y2="{H-10}" stroke="{LINE}"/>')
    for i, (lab, v) in enumerate(rows):
        y = T + i * rowh; w = span * min(abs(v), 1.0)
        x = mid if v >= 0 else mid - w
        out.append(f'<text x="{L-10}" y="{y+21}" font-size="12" text-anchor="end" fill="{INK}" font-family="{SANS}">{lab}</text>')
        out.append(f'<rect x="{x:.1f}" y="{y+8}" width="{w:.1f}" height="18" rx="3" fill="{ACC if v>=0 else ACC3}"/>')
        tx = mid + w + 6 if v >= 0 else mid - w - 6
        out.append(f'<text x="{tx:.1f}" y="{y+21}" font-size="12" text-anchor="{"start" if v>=0 else "end"}" fill="{INK}" font-family="{MONO}">{v*100:.0f}%</text>')
    out.append("</svg>"); return "\n".join(out)

for name, fn in [("chart-sizing", chart_sizing), ("chart-pipeline", chart_pipeline), ("chart-tco", chart_tco), ("chart-matrix", chart_matrix), ("chart-margins", chart_margins)]:
    open(HERE + name + ".svg", "w").write(fn())

prog_rows = "\n".join(f'<tr><td><b>{p["name"]}</b></td><td>{p["quad"]}</td><td class="num">{p["mw"]:,.0f}</td><td class="num">{p["yr"]}</td><td class="num">{p["mw26"]:,.0f}</td><td class="num">{p["gpus26"]:,.0f}</td><td class="num">${p["capex26"]/1000:,.1f}B</td></tr>' for p in progs)
quad_rows = "\n".join(f'<tr><td><b>{q["name"]}</b></td><td class="num">{q["mw"]:,.0f}</td><td class="num">{q["mw26"]:,.0f}</td><td class="num">{q["gpus26"]:,.0f}</td><td class="num">${q["capex26"]/1000:,.1f}B</td><td class="num">${q["gpusys"]/1000:,.1f}B</td></tr>' for q in quads)
ctry = ["United States", "Saudi Arabia", "South Korea", "Germany", "India"]
tco_rows = "".join(f'<tr><td>{c}</td><td class="num">${G["G7"][i]:,.2f}B</td><td class="num">{G["G14"][i]*100:.0f}%</td><td class="num">${G["G11"][i]:,.2f}B</td><td class="num">${G["G13"][i]:,.1f}B</td><td class="num">${G["G18"][i]:,.2f}</td><td class="num">${G["G18c"][i]:,.2f}</td></tr>' for i, c in enumerate(ctry))

tmpl = open(HERE + "index.template.html").read()
rep = {"PROG_ROWS": prog_rows, "QUAD_ROWS": quad_rows, "TCO_ROWS": tco_rows,
       "CHART_SIZING": chart_sizing(), "CHART_PIPELINE": chart_pipeline(), "CHART_TCO": chart_tco(), "CHART_MATRIX": chart_matrix(), "CHART_MARGINS": chart_margins()}
for k in V:
    v = V[k]
    if isinstance(v, (int, float)):
        rep[k] = v
# formatted helpers
F = {}
F["td_A3"] = bn(V["td_A3"], 0); F["td_A5"] = bn(V["td_A5"], 0); F["td_A7"] = bn(V["td_A7"], 0); F["td_A8"] = gw(V["td_A8"])
F["td_B4"] = bn(V["td_B4"], 1); F["td_B5"] = bn(V["td_B5"], 1); F["td_B7"] = bn(V["td_B7"], 1)
F["td_C1"] = bn(V["td_C1"], 1); F["td_C2"] = bn(V["td_C2"], 1); F["td_C3"] = bn(V["td_C3"], 1); F["td_C4"] = bn(V["td_C4"], 1); F["td_C5"] = bn(V["td_C5"], 1); F["td_C7"] = bn(V["td_C7"], 0)
F["td_R1"] = bn(V["td_R1"], 0); F["td_R2"] = bn(V["td_R2"], 0); F["td_R3"] = bn(V["td_R3"], 0); F["td_R4"] = bn(V["td_R4"], 0); F["td_R5"] = pct(V["td_R5"]); F["td_R6"] = bn(V["td_R6"], 0); F["td_R7"] = pct(V["td_R7"]); F["td_R9"] = f"{V['td_R9']:.1f}x"; F["td_R10"] = pct(V["td_R10"], 1); F["td_R11"] = pct(V["td_R11"])
F["bu_gw26"] = gw(V["bu_gw26"]); F["bu_capex26"] = bn(V["bu_capex26"], 0); F["bu_gpu26"] = bn(V["bu_gpu26"], 0); F["bu_gpus26"] = num(V["bu_gpus26"]); F["bu_pipe_gw"] = gw(V["bu_pipe_gw"], 0); F["bu_pipe_capex"] = bn(V["bu_pipe_capex"], 0); F["bu_pipe_years"] = f"{V['bu_pipe_years']:.0f}"
F["gulf_share"] = pct(V["gulf_share"]); F["eu_ann_share"] = pct(V["eu_ann_share"]); F["eu_en_share"] = pct(V["eu_en_share"])
for i, c in enumerate(["us", "sa", "kr", "de", "in"]):
    F[f"g11_{c}"] = f"${G['G11'][i]:,.2f}B"; F[f"g17_{c}"] = usd(G["G17"][i]); F[f"g18_{c}"] = usd(G["G18"][i]); F[f"g18c_{c}"] = usd(G["G18c"][i]); F[f"g14_{c}"] = pct(G["G14"][i]); F[f"g13_{c}"] = f"${G['G13'][i]:,.1f}B"; F[f"g7_{c}"] = f"${G['G7'][i]:,.2f}B"
F["g15"] = num(G["G15"][0]); F["g21"] = usd(G["G21"][0], 0)
for code, vals in KV.items():
    for i, v in enumerate(vals, start=1):
        if isinstance(v, (int, float)):
            if code in ("K4",): F[f"k_{code}_{i}"] = num(v)
            elif code in ("K11", "K16", "K18"): F[f"k_{code}_{i}"] = f"${v:,.2f}B"
            elif code in ("K19",): F[f"k_{code}_{i}"] = pct(v)
            elif code in ("K21", "K21b", "K22", "K23"): F[f"k_{code}_{i}"] = usd(v)
            elif code in ("K24",): F[f"k_{code}_{i}"] = (f"${v/1000:,.1f}B" if abs(v) >= 1000 else f"${v:,.0f}M")
        else:
            F[f"k_{code}_{i}"] = str(v)
for k in ["P1", "P2", "P3", "P4", "P5", "P5b", "P5c", "P5d", "P6", "P7", "P8", "P9", "P10", "P11", "P12", "P13", "P14", "P15", "P16", "P17", "P18", "P21", "P22", "P26", "P27", "P28"]:
    F["p_" + k] = usd(V["p_" + k], 2 if V["p_" + k] < 1000 else 0)
F["p_P19"] = pct(V["p_P19"]); F["p_P20"] = f"{V['p_P20']:.1f}x"; F["p_P23"] = f"${V['p_P23']:,.0f}M"; F["p_P24"] = pct(V["p_P24"]); F["p_P25"] = pct(V["p_P25"]); F["p_P25b"] = pct(V["p_P25b"]); F["p_P25c"] = pct(V["p_P25c"]); F["p_P29"] = pct(V["p_P29"])
F["v_VA3"] = bn(V["v_VA3"], 0); F["v_VA4"] = bn(V["v_VA4"], 0); F["v_VA5"] = f"{V['v_VA5']:.1f}"; F["v_VA6"] = bn(V["v_VA6"], 0); F["v_VA7"] = bn(V["v_VA7"], 0); F["v_VA8"] = f"{V['v_VA8']*12:.1f} months"; F["v_VA9"] = bn(V["v_VA9"], 0); F["v_VA10"] = bn(V["v_VA10"], 1); F["v_VA11"] = f"{V['v_VA11']*365:.0f} days"; F["v_VA12"] = bn(V["v_VA12"], 0); F["v_VA13"] = bn(V["v_VA13"], 0); F["v_VA14"] = f"{V['v_VA14']:.1f}"
F["v_VB1"] = f"{V['v_VB1']:.0f}x"; F["v_VB3"] = pct(V["v_VB3"]); F["v_VC1"] = f"${V['v_VC1']:,.1f}B"; F["v_VC2"] = f"${V['v_VC2']:,.1f}B"; F["v_VC3"] = pct(V["v_VC3"]); F["v_VC3b"] = pct(V["v_VC3b"]); F["v_VC4"] = f"${V['v_VC4']:,.1f}B"; F["v_VC5"] = f"${V['v_VC5']:,.1f}B"; F["v_VC6"] = pct(V["v_VC6"]); F["v_VC6b"] = pct(V["v_VC6b"]); F["v_VC6b_abs"] = pct(abs(V["v_VC6b"])); F["v_VC7"] = usd(V["v_VC7"])
rep.update(F)
out = tmpl
for k, v in rep.items():
    out = out.replace("{{" + k + "}}", str(v))
left = sorted(set(re.findall(r"\{\{[a-zA-Z0-9_]+\}\}", out)))
if left:
    raise SystemExit("unfilled: " + ", ".join(left))
open(HERE + "index.html", "w").write(out)
print("wrote index.html", len(out))
